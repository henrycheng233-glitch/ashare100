# -*- coding: utf-8 -*-
"""
从集成人提供的《全部A股-10.21市值前一百.xlsx》生成 stocks.py。

为什么要重新生成：东财 screener 的市值是总市值（含 H 股），集成人的表是 A 股口径。
两者在第 97–100 名边界不一致：
  东财口径进前 100、表中不在：601336 新华保险（表内 106）、601995 中金公司（129）
  表中进前 100、东财口径不在：300274 阳光电源（97）、300033 同花顺（100）
以集成人的表为准。

名称取自东财（market_universe.json 的 abbr），因为 Excel 的证券名称列带 XD/XR 前缀且被截断
（"XD工商银" 实际是 "XD工商银行" 截断），不适合当 name_abbr。
"""
import json
import os
import re

import openpyxl

XLSX = r"C:/Users/cuiya/Documents/xwechat_files/yangspecialster_4725/msg/file/2026-10/全部A股-10.21市值前一百.xlsx"
HERE = os.path.dirname(os.path.abspath(__file__))


def exchange_of(code):
    if code.startswith("6"):
        return "SH"
    if code.startswith(("0", "3")):
        return "SZ"
    if code.startswith(("4", "8")):
        return "BJ"
    raise ValueError(code)


def main():
    ws = openpyxl.load_workbook(XLSX, data_only=True)["全部A股"]
    rows = []
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not r[0]:
            continue
        try:
            mv = float(r[4])
        except (TypeError, ValueError):
            continue
        rows.append({
            "key": str(r[0]).strip(),
            "name_xlsx": str(r[1]).strip(),
            "full": str(r[3]).strip(),
            "mv": mv,
        })

    top = rows[:100]
    assert len(top) == 100, f"前 100 行只有 {len(top)} 条"

    # 东财简称，用于清洗 Excel 截断名
    uni = {}
    p = os.path.join(HERE, "market_universe.json")
    if os.path.exists(p):
        for x in json.load(open(p, encoding="utf-8")):
            uni[x["code"]] = x.get("abbr")

    out_lines = []
    for i, x in enumerate(top, 1):
        code = x["key"].split(".")[0]
        ex = x["key"].split(".")[1]
        assert exchange_of(code) == ex, f"{code} 交易所不符 {ex}"
        name = uni.get(code) or re.sub(r"^(XD|XR|DR)", "", x["name_xlsx"])
        out_lines.append(
            f'    ("{code}", "{name}", {i}),'
        )

    body = "\n".join(out_lines)
    content = f'''# -*- coding: utf-8 -*-
"""C 模块：A 股市值前 100 名单。

来源：集成人《全部A股-10.21市值前一百.xlsx》前 100 行（A 股口径，按总市值降序）。
由 build_stocks.py 生成，不要手改——名单变更请改 Excel 后重新生成。

主键统一 `代码.交易所`（规范 §4.1）。
"""
from collections import Counter

# (代码, 简称, 排名)
_RAW = [
{body}
]


def exchange_of(code):
    if code.startswith("6"):
        return "SH"
    if code.startswith(("0", "3")):
        return "SZ"
    if code.startswith(("4", "8")):
        return "BJ"
    raise ValueError(f"无法判定交易所: {{code}}")


STOCKS = [
    {{
        "code": code,
        "exchange": exchange_of(code),
        "key": f"{{code}}.{{exchange_of(code)}}",
        "name": name,
        "rank": rank,
    }}
    for code, name, rank in _RAW
]

BY_CODE = {{s["code"]: s for s in STOCKS}}
KEYS = [s["key"] for s in STOCKS]


if __name__ == "__main__":
    print(f"共 {{len(STOCKS)}} 家")
    print("交易所分布:", dict(Counter(s["exchange"] for s in STOCKS)))
    dup = [k for k, v in Counter(KEYS).items() if v > 1]
    print("重复主键:", dup or "无")
    print("首 3:", [s["key"] for s in STOCKS[:3]])
    print("末 3:", [s["key"] for s in STOCKS[-3:]])
'''
    open(os.path.join(HERE, "stocks.py"), "w", encoding="utf-8", newline="\n").write(content)

    print(f"已生成 stocks.py，{len(top)} 家")
    print("首 3:", [(x['key'], uni.get(x['key'].split('.')[0]) or x['name_xlsx']) for x in top[:3]])
    print("末 3:", [(x['key'], uni.get(x['key'].split('.')[0]) or x['name_xlsx']) for x in top[-3:]])
    print("新增 2 家:", [x['key'] for x in top if x['key'] in ('300033.SZ', '300274.SZ')])
    print("移除 2 家:", [k for k in ('601336.SH', '601995.SH') if k not in {x['key'] for x in top}])


if __name__ == "__main__":
    main()
